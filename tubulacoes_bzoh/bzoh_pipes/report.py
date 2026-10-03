"""Calculation report (memorial de cálculo) and line list outputs."""

import csv
import datetime

import numpy as np

from . import hydraulics as hy
from .lines import CW_T_IN, CW_T_OUT, DESIGN_FACTOR, TURNDOWN
from .properties import pure_viscosity
from .sizing import FR_MAX, NPSH_MARGIN, NPSH_REQUIRED
from .tables import INCH, PSI_100FT, SPECS, economic_velocity

GROUPS = ["Processo (PFD)", "Colunas (auxiliares)", "Utilidades"]
SERVICE_SHORT = {
    "discharge": "Recalque", "suction": "Sucção (saturado)",
    "suction_vac": "Sucção (vácuo)", "reboiler_feed": "Alim. refervedor",
    "gravity": "Gravidade", "vapor": "Vapor", "vapor_vac": "Vapor (vácuo)",
    "two_phase": "Bifásico", "cw": "Água de resfriamento", "steam": "Vapor d'água",
    "condensate": "Condensado",
}
PHASE = {"L": "Líquido", "V": "Vapor", "LV": "Bifásico"}
FLUID_CODES = [
    ("P", "processo - líquido"), ("PV", "processo - vapor"),
    ("PM", "processo - bifásico"), ("AR", "água de resfriamento"),
    ("VB / VM / VA", "vapor d'água de baixa / média / alta pressão"),
    ("CD", "condensado de vapor"),
]
SHORT_GOV = {
    'DN mínimo da especificação (3/4")': 'DN mínimo',
    'Velocidade econômica': 'Velocidade econômica',
    'Perda de carga econômica (ΔP/100 m)': 'ΔP econômico',
    'ΔP admissível em vácuo (1 % de P)': 'ΔP em vácuo',
    'Autoventilação (Simpson, Fr ≤ 0,31)': 'Autoventilação',
    'Velocidade máxima admissível': 'Velocidade máxima',
    'Velocidade erosional (API RP 14E)': 'Vel. erosional',
}

# APEA (Aspen Process Economic Analyzer) line diameters for the same services
APEA_DN = {
    "111": 20, "114": 20, "115": 20, "116": 20,
    "121": 20, "124": 20, "125": 20, "126": 20,
    "131": 50, "134": 50, "135": 50, "136": 50,
    "141": 50, "144": 50, "145": 50, "146": 50,
    "151": 20, "154": 20, "155": 20, "156": 20,
    "301": 20, "302": 50, "303": 20, "304": 50, "305": 20,
    "203": 50, "205": 25, "207": 50,
}


def br(x, nd=2):
    """Brazilian number format (1.234,56)."""
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "—"
    s = f"{x:,.{nd}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def sig(x, n=3):
    """Brazilian format with n significant digits (no thousands separator)."""
    if x is None or not np.isfinite(x) or x == 0:
        return "0" if x == 0 else "—"
    nd = max(0, n - 1 - int(np.floor(np.log10(abs(x)))))
    return f"{x:.{nd}f}".replace(".", ",")


def msig(x, n=3):
    """sig() for use inside LaTeX math (decimal comma as {,})."""
    return sig(x, n).replace(",", "{,}")


def mbr(x, nd=2):
    """br() for use inside LaTeX math."""
    return br(x, nd).replace(",", "{,}").replace(".", "\\,")


def table(headers, rows, align=None):
    align = align or ["---"] * len(headers)
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(align) + " |"]
    out += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def route(ln):
    return f"{ln.origin} → {ln.destination}"


def nps_dn(s):
    return f'{s.label} (DN {s.DN})'


# --------------------------------------------------------------------------
# Line list files
# --------------------------------------------------------------------------
LIST_COLUMNS = [
    "Tag", "Linha", "Grupo", "Código do fluido", "Descrição", "De", "Para", "Fase",
    "Serviço", "Classe", "Material", "Norma do tubo", "Schedule", "Flanges",
    "NPS", "DN", "DE [mm]", "Espessura [mm]", "DI [mm]", "Vazão mássica [kg/h]",
    "Q normal [m³/h]", "Q projeto [m³/h]", "T [°C]", "P [kPa abs]",
    "ρ [kg/m³]", "μ [mPa·s]", "D econômico [mm]", "v projeto [m/s]",
    "v econômica [m/s]", "Re", "ΔP [kPa/100 m]", "ΔP linha [kPa]",
    "Critério governante", "Observações", "Fonte dos dados",
]


def _row(r):
    ln, s, c = r.line, r.chosen.size, r.chosen
    return [
        r.tag, ln.number, ln.group, ln.code, ln.description, ln.origin,
        ln.destination, PHASE[ln.phase], r.service.label, r.spec.code,
        r.spec.material, r.spec.pipe_std, s.schedule, r.spec.rating, s.label, s.DN,
        s.OD * 1e3, s.wall * 1e3, s.ID * 1e3, ln.mass_kg_h, ln.Q_normal * 3600,
        ln.Q_design * 3600, ln.T_C, ln.P_kPa, ln.rho, ln.mu * 1e3, r.D_ec * 1e3,
        c.v_d, c.v_ec, c.Re, c.dpdl / 10.0, r.dP_line / 1e3, r.governing,
        " ".join(r.notes), ln.source,
    ]


def write_csv(results, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(LIST_COLUMNS)
        for r in results:
            w.writerow([(f"{v:.6g}".replace(".", ",") if isinstance(v, float) else v)
                        for v in _row(r)])


def write_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Lista de linhas"
    ws.append(LIST_COLUMNS)
    head = PatternFill("solid", fgColor="1F3B57")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in results:
        ws.append([float(v) if isinstance(v, (float, np.floating)) else v
                   for v in _row(r)])
    fmts = {"DE [mm]": "0.00", "Espessura [mm]": "0.00", "DI [mm]": "0.00",
            "Vazão mássica [kg/h]": "0.00", "Q normal [m³/h]": "0.000",
            "Q projeto [m³/h]": "0.000", "T [°C]": "0.0", "P [kPa abs]": "0.00",
            "ρ [kg/m³]": "0.000", "μ [mPa·s]": "0.0000", "D econômico [mm]": "0.0",
            "v projeto [m/s]": "0.00", "v econômica [m/s]": "0.00", "Re": "0",
            "ΔP [kPa/100 m]": "0.000", "ΔP linha [kPa]": "0.000"}
    widths = {"Tag": 22, "Descrição": 46, "De": 26, "Para": 28, "Serviço": 40,
              "Material": 30, "Norma do tubo": 30, "Schedule": 10, "Flanges": 36,
              "Critério governante": 36, "Observações": 60, "Fonte dos dados": 28,
              "Grupo": 20}
    for j, name in enumerate(LIST_COLUMNS, start=1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = widths.get(name, 13)
        if name in fmts:
            for i in range(2, ws.max_row + 1):
                ws[f"{col}{i}"].number_format = fmts[name]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32

    ws2 = wb.create_sheet("Classes de material")
    ws2.append(["Classe", "Material", "Tubo", "Schedules", "Flanges", "Serviço",
                "P projeto [bar g]"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    for sp in SPECS.values():
        ws2.append([sp.code, sp.material, sp.pipe_std, sp.schedules, sp.rating,
                    sp.service, sp.P_design / 1e5])
    for j, wdt in enumerate([8, 30, 34, 52, 40, 60, 16], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = wdt
    wb.save(path)


# --------------------------------------------------------------------------
# Memorial
# --------------------------------------------------------------------------
def _gov_short(r):
    return SHORT_GOV.get(r.governing, r.governing)


def _result_table(results):
    rows = []
    for r in results:
        s = r.chosen.size
        rows.append([f"`{r.tag}`", r.line.description, route(r.line),
                     PHASE[r.line.phase], f"{r.spec.code} - {r.spec.material}",
                     f"**{s.label}**", s.schedule, br(s.ID * 1e3, 1),
                     sig(r.chosen.v_d, 2), _gov_short(r)])
    return table(["Tag (P&ID)", "Serviço", "De → Para", "Fase", "Material (classe)",
                  "NPS", "Schedule", "DI [mm]", "v proj. [m/s]", "Critério"], rows,
                 ["---", "---", "---", "---", "---", ":---:", ":---:", "---:",
                  "---:", "---"])


def _calc_table(results):
    rows = []
    for r in results:
        ln, c, s = r.line, r.chosen, r.chosen.size
        extra = ""
        if r.service.rule == "simpson":
            extra = f"Fr = {sig(hy.simpson_froude(ln.Q_design, s.ID), 2)}"
        elif r.service.rule == "two_phase":
            extra = f"Ve = {sig(hy.erosional_velocity(ln.rho), 3)}"
        else:
            extra = sig(c.v_ec, 3)
        rows.append([ln.number, SERVICE_SHORT[ln.service],
                     sig(ln.Q_normal * 3600, 3), sig(ln.Q_design * 3600, 3),
                     sig(ln.rho, 4), sig(ln.mu * 1e3, 3), sig(r.D_ec * 1e3, 3),
                     f"{s.label} / {br(s.ID * 1e3, 1)}", sig(c.v_d, 3), extra,
                     br(c.Re, 0), sig(c.dpdl / 10.0, 3), _gov_short(r)])
    return table(["Linha", "Serviço", "Qₙ [m³/h]", "Q_proj [m³/h]", "ρ [kg/m³]",
                  "μ [mPa·s]", "D_ec [mm]", "NPS / DI [mm]", "v_proj [m/s]",
                  "v_ec [m/s] (ou Fr, Ve)", "Re", "ΔP [kPa/100 m]", "Critério"], rows,
                 ["---", "---"] + ["---:"] * 10 + ["---"])


def _example(r, title):
    """Step-by-step calculation of one line."""
    ln, svc, spec = r.line, r.service, r.spec
    c = r.chosen
    out = [f"### {title}", ""]
    out.append(f"**Linha {ln.number} — {ln.description}** ({route(ln)}); "
               f"serviço: *{svc.label}*; classe **{spec.code}** ({spec.material}).")
    out.append("")
    data = [["Vazão mássica", f"{br(ln.mass_kg_h, 2)} kg/h"],
            ["Temperatura / pressão", f"{br(ln.T_C, 1)} °C / {br(ln.P_kPa, 2)} kPa abs"],
            ["Massa específica ρ", f"{sig(ln.rho, 4)} kg/m³"],
            ["Viscosidade μ", f"{sig(ln.mu * 1e3, 3)} mPa·s"],
            ["Vazão volumétrica normal Qₙ", f"{sig(ln.Q_normal, 4)} m³/s = {sig(ln.Q_normal * 3600, 4)} m³/h"],
            ["Vazão de projeto Q = 1,30 Qₙ", f"{sig(ln.Q_design, 4)} m³/s"]]
    out.append(table(["Dado", "Valor"], data))
    out.append("")
    if svc.rule in ("discharge", "suction", "gas"):
        rule_txt = {"discharge": r"v_{ec} = (5 + D/3)\ \text{ft/s}",
                    "suction": r"v_{ec} = (1{,}3 + D/6)\ \text{ft/s}",
                    "gas": r"v_{ec} = 20\,D\ \text{ft/s}"}[svc.rule]
        out.append("**Passo 1 — diâmetro econômico.** Com a heurística de Walas/Turton "
                   f"$${rule_txt}\\quad (D\\ \\text{{em polegadas}})$$ "
                   "resolve-se, por iteração, $Q = \\dfrac{\\pi D^2}{4}\\,v_{ec}(D)$:")
        out.append("")
        out.append(f"$$D_{{ec}} = {msig(r.D_ec * 1e3, 4)}\\ \\text{{mm}} "
                   f"\\;({msig(r.D_ec / INCH, 3)}\\ \\text{{in}}),\\qquad "
                   f"v_{{ec}}(D_{{ec}}) = {msig(economic_velocity(svc.rule, r.D_ec), 3)}\\ \\text{{m/s}}$$")
    elif svc.rule == "simpson":
        out.append("**Passo 1 — diâmetro mínimo autoventilado** (Simpson, 1968): "
                   "$Fr = v/\\sqrt{gD} \\le 0{,}31 \\Rightarrow "
                   "D_{min} = \\left(\\dfrac{4Q}{\\pi\\sqrt{g}\\,Fr_{max}}\\right)^{0{,}4}$ "
                   f"$= {msig(r.D_ec * 1e3, 4)}$ mm.")
    out.append("")
    i = r.candidates.index(c)
    show = r.candidates[max(0, i - 2): i + 2]
    rows = []
    for cd in show:
        crit = []
        if svc.rule in ("discharge", "suction", "gas"):
            crit.append(sig(cd.v_ec, 3))
        if svc.dp_frac_P:
            crit.append(f"{br(cd.dpdl * ln.L_eq, 1)} Pa (lim. {br(svc.dp_frac_P * ln.P_kPa * 1e3, 1)})")
        elif svc.rule != "simpson":
            crit.append(sig(cd.dpdl / 10.0, 3))
        if svc.rule == "simpson":
            crit.append(sig(hy.simpson_froude(ln.Q_design, cd.size.ID), 3))
        verdict = "✔ adotado" if cd is c else ("✔" if not cd.fails else "✘ " + ", ".join(cd.fails))
        rows.append([f"{cd.size.label} {cd.size.schedule}", br(cd.size.ID * 1e3, 2),
                     sig(cd.v_d, 3)] + crit + [verdict])
    heads = ["NPS", "DI [mm]", "v_proj [m/s]"]
    if svc.rule in ("discharge", "suction", "gas"):
        heads.append("v_ec(DI) [m/s]")
    if svc.dp_frac_P:
        heads.append(f"ΔP linha (L_eq = {br(ln.L_eq, 0)} m)")
    elif svc.rule != "simpson":
        heads.append(f"ΔP [kPa/100 m] (lim. {sig(svc.dp_limit / 10, 3)})")
    if svc.rule == "simpson":
        heads.append("Fr (lim. 0,31)")
    heads.append("Verificação")
    out.append("**Passo 2 — varredura dos diâmetros comerciais** "
               f"({spec.pipe_std.split(' (')[0]}; {spec.schedules}):")
    out.append("")
    out.append(table(heads, rows))
    out.append("")
    s = c.size
    eps = spec.roughness
    out.append("**Passo 3 — verificação hidráulica do diâmetro adotado** "
               f"({s.label} {s.schedule}, DI = {br(s.ID * 1e3, 2)} mm):")
    out.append("")
    out.append(f"- Área: $A = \\pi D^2/4 = {msig(np.pi * s.ID ** 2 / 4, 4)}$ m²; "
               f"$v = Q/A = {msig(c.v_d, 3)}$ m/s (normal: {sig(c.v_n, 3)} m/s).")
    out.append(f"- Reynolds: $Re = \\rho v D/\\mu = {mbr(c.Re, 0)}$ "
               f"({'laminar' if c.Re < 2300 else 'turbulento'}); "
               f"$\\varepsilon/D = {msig(eps / s.ID, 3)}$.")
    out.append(f"- Fator de atrito ({'64/Re' if c.Re < 2300 else 'Colebrook-White'}): "
               f"$f = {msig(c.f, 4)}$.")
    out.append(f"- Darcy-Weisbach: $\\Delta P/L = f\\,\\dfrac{{1}}{{D}}\\,\\dfrac{{\\rho v^2}}{{2}} "
               f"= {msig(c.dpdl, 4)}$ Pa/m $= {msig(c.dpdl / 10, 3)}$ kPa/100 m; "
               f"na linha ($L_{{eq}}$ = {br(ln.L_eq, 0)} m): ΔP = {sig(r.dP_line / 1e3, 3)} kPa.")
    w = r.wall
    out.append(f"- Espessura (ASME B31.3, P = {br(w['P_design'] / 1e5, 0)} bar g): "
               f"$t_{{min}} = {mbr(w['t_min'] * 1e3, 2)}$ mm ≤ $t_{{nom}} = {mbr(w['t_nom'] * 1e3, 2)}$ mm ✔"
               + (f"; vácuo: $P_{{cr}}/3 = {mbr(w['P_cr'] / 3e6, 2)}$ MPa ≥ 0,101 MPa ✔"
                  if "P_cr" in w else "") + ".")
    out.append("")
    out.append(f"**Resultado:** `{r.tag}` — critério governante: *{r.governing}*.")
    out.append("")
    return "\n".join(out)


def _apea_table(results):
    by = {r.line.number: r for r in results}
    rows = []
    for num, dn in APEA_DN.items():
        r = by.get(num)
        if r is None:
            continue
        s = r.chosen.size
        A = np.pi * (dn / 1e3) ** 2 / 4
        v_apea = r.line.Q_design / A
        rows.append([num, r.line.description, f"DN {dn}", f"{s.label} (DN {s.DN})",
                     sig(v_apea, 3), sig(r.chosen.v_d, 3)])
    return table(["Linha", "Serviço", "APEA", "Este memorial",
                  "v no DN do APEA [m/s]", "v adotada [m/s]"], rows)


def _econ_table(results):
    rows = []
    from .sizing import nearest_preferred
    for r in results:
        if r.governing.startswith("DN mínimo") or r.service.rule in ("simpson",):
            continue
        if r.chosen.Re < 4000:
            continue
        e = r.econ_diameters
        corr = "Sinnott (aço inox)" if r.spec.code.startswith("S") else "Sinnott (aço carbono)"
        rows.append([r.line.number, r.line.description,
                     f"{r.chosen.size.label}",
                     f"{sig(r.D_ec * 1e3, 3)}",
                     f"{sig(e['Peters-Timmerhaus'] * 1e3, 3)} → {nearest_preferred(r.spec.code, e['Peters-Timmerhaus']).label}",
                     f"{sig(e[corr] * 1e3, 3)} → {nearest_preferred(r.spec.code, e[corr]).label}"])
    return table(["Linha", "Serviço", "NPS adotado", "D_ec Walas [mm]",
                  "Peters–Timmerhaus [mm] → NPS", "Sinnott [mm] → NPS"], rows)


def write_memorial(results, sens, extra, path, fig_paths):
    today = datetime.date.today().strftime("%d/%m/%Y")
    n = len(results)
    by_group = {g: [r for r in results if r.line.group == g] for g in GROUPS}
    by_spec = {k: sum(1 for r in results if r.spec.code == k) for k in SPECS}
    dmin = sum(1 for r in results if r.governing.startswith("DN mínimo"))
    big = sorted(results, key=lambda r: -r.chosen.size.ID)[:6]
    L = []
    a = L.append

    a("# Memorial de Cálculo — Dimensionamento das Tubulações do P&ID")
    a("")
    a("**Planta de álcool benzílico — Área 2 (hidrólise do acetato de benzila) · "
      "Projeto Base · 2,37 t/d**")
    a("")
    a(table(["Item", "Descrição"], [
        ["Documento", "MC-TUB-001 · Revisão 0"],
        ["Disciplina", "Projetos de Indústrias Químicas I — UNESP, Instituto de Química (Araraquara)"],
        ["Base de dados", "Aspen Plus V15 (Projeto Base, NRTL-HOC) e relatório PIQ I — "
                          "Tabelas 4 a 7, 26, 29 e 30; Anexos 7, 9, 10, 13 e 16"],
        ["Método", "Velocidade (vazão) econômica — heurísticas de Walas (1988) e "
                   "Turton et al. (2018), com verificação hidráulica e mecânica"],
        ["Arquivos gerados", "`resultados/lista_de_linhas.xlsx` e `.csv` (lista de linhas "
                             "completa), figuras em `resultados/`"],
        ["Gerado em", f"{today} por `tubulacoes_bzoh/run_dimensionamento.py`"],
    ]))
    a("")
    a("## Resultado")
    a("")
    a(f"Foram dimensionadas **{n} linhas**: {len(by_group[GROUPS[0]])} correntes do PFD, "
      f"{len(by_group[GROUPS[1]])} linhas auxiliares das colunas T-101 a T-105 (vapor de "
      "topo, condensado, sucção e recalque de refluxo, alimentação e retorno dos "
      f"refervedores) e {len(by_group[GROUPS[2]])} linhas de utilidades (água de "
      "resfriamento, vapor e condensado).")
    a("")
    a(f"- **Materiais:** inox **AISI 316L** (classe S1) em {by_spec['S1']} linhas de processo, "
      "por causa do ácido acético e do meio hidrolítico; inox **AISI 304L** (S2) nas "
      f"{by_spec['S2']} alimentações limpas a 25 °C (água e acetato de benzila); "
      f"**aço carbono ASTM A106 Gr. B** nas {by_spec['C1'] + by_spec['C2']} linhas de "
      f"utilidades (C1, Classe 150: {by_spec['C1']} linhas; C2, Classe 600, vapor de "
      f"alta: {by_spec['C2']} linhas).")
    a(f"- **Diâmetros:** de 3/4\" a {big[0].chosen.size.label}. Como a planta é pequena "
      "(vazões de líquido de processo abaixo de 1 m³/h), "
      f"**{dmin} linhas ficam no diâmetro mínimo de 3/4\"**: o diâmetro econômico "
      "delas é menor que 3/4\", e o diâmetro mínimo da especificação é que governa.")
    a("- **Linhas de maior diâmetro:** "
      + "; ".join(f"`{r.tag}` ({r.line.description})" for r in big) + ".")
    a("")
    a("**Formato do tag:** `NPS-código do fluido-número da linha-classe de material`. "
      "Os números 001 a 034 são as correntes do PFD (mesma numeração do Aspen), "
      "1*k*1 a 1*k*6 são as linhas auxiliares da coluna T-10*k*, 2xx é água de "
      "resfriamento, 30*k* é vapor e 40*k* é condensado do refervedor da T-10*k*. "
      "Códigos de fluido: " + "; ".join(f"**{c}** = {d}" for c, d in FLUID_CODES) + ".")
    a("")
    for g in GROUPS:
        a(f"### Lista de linhas — {g}")
        a("")
        a(_result_table(by_group[g]))
        a("")
    a("Classes de material (detalhes na Seção 6):")
    a("")
    a(table(["Classe", "Material", "Tubo", "Schedule", "Flanges", "Aplicação"],
            [[sp.code, sp.material, sp.pipe_std, sp.schedules, sp.rating, sp.service]
             for sp in SPECS.values()]))
    a("")

    # 1 ---------------------------------------------------------------
    a("## 1. Objetivo e escopo")
    a("")
    a("Este memorial define o **diâmetro nominal, o schedule e o material** de todas as "
      "linhas que aparecem no P&ID da Área 2: as correntes do PFD (Figuras 1 e 2 do "
      "relatório), as linhas internas ao sistema de cada coluna, que o PFD do Aspen "
      "não mostra, mas o P&ID precisa (condensador, vaso e bomba de refluxo, refervedor "
      "kettle), e as conexões de utilidades. O dimensionamento é feito pelo **método da "
      "velocidade (vazão) econômica**, verificado por perda de carga, por critérios "
      "específicos de cada serviço (vácuo, gravidade, bifásico), por NPSH e pela "
      "espessura de parede (ASME B31.3). Ficam fora do escopo o sistema de vácuo "
      "(ejetores ou bombas de vácuo, que ainda não estão orçados no relatório), as "
      "linhas de alívio (PSV) e o detalhamento isométrico.")
    a("")

    # 2 ---------------------------------------------------------------
    a("## 2. Dados de entrada")
    a("")
    a(table(["Dado", "Fonte"], [
        ["Vazões, T, P, fase, massa específica, massa molar e composição das 34 correntes",
         "Aspen Plus — *Results Summary/Streams* (planilha `BZOH_BASE_CORRENTES.xlsx`, "
         "extraída para `dados/correntes_aspen.csv`)"],
        ["Destilado, refluxo, fundo, boil-up, T de topo/fundo e cargas do condensador e "
         "do refervedor", "Relatório PIQ I — Tabela 7 e Anexos 7, 9, 10, 13 e 16 "
         "(`dados/colunas.csv`)"],
        ["Pressão de operação das colunas", "Tabela 26 (T-101: 50 torr; T-105: 100 torr; "
         "demais: 1 atm)"],
        ["Cargas dos resfriadores E-101 a E-103 e do reator R-101", "Tabela 6 (`dados/resfriadores.csv`)"],
        ["Níveis de vapor (VB, VM e VA) e água de resfriamento", "Tabela 6, utilidades do "
         "Aspen/APEA e tabelas de vapor (`dados/utilidades.csv`)"],
        ["Materiais dos equipamentos", "Tabela 29 e Seção 4"],
        ["Cenários de produção (base, +30 % e −30 %)", "Tabela 34"],
    ]))
    a("")
    a("A tabela de correntes do Aspen não traz viscosidade; ela foi estimada "
      "(Seção 5), e a Seção 9.5 mostra que o resultado é pouco sensível a essa "
      "estimativa.")
    a("")

    # 3 ---------------------------------------------------------------
    a("## 3. Premissas de projeto")
    a("")
    a(table(["Premissa", "Valor adotado", "Justificativa"], [
        ["Vazão de projeto", f"Q = {br(DESIGN_FACTOR, 2)} × Qₙ",
         "Cobre o cenário de produção 30 % acima (Tabela 34), sem trocar tubulação"],
        ["Vazão mínima de verificação", f"{br(TURNDOWN, 2)} × Qₙ",
         "Cenário 30 % abaixo (golfadas em linhas bifásicas)"],
        ["Diâmetro mínimo", '3/4" (DN 20) em todas as classes',
         "Resistência mecânica e suportação em rack; mesmo mínimo do APEA e dos bocais "
         "da T-101 (Tabela 19)"],
        ["Diâmetros não usados", '1.1/4", 3.1/2" e 5"', "Pouco comerciais (Telles)"],
        ["Comprimento equivalente (trecho reto + acessórios)",
         "30 m processo; 25 m refluxo; 20 m vapor de topo; 5 m linhas junto à coluna; "
         "40 m AR; 30 m vapor; 10 m condensado",
         "Estimativa sem isométrico; só afeta a ΔP total da linha e o critério de vácuo"],
        ["Água de resfriamento", f"{br(CW_T_IN, 0)} → {br(CW_T_OUT, 0)} °C; "
         "cp = 4,18 kJ/(kg·K)", "Mesma faixa do condensador da T-101 (Tabela 23)"],
        ["Vapor (saturado)", "VB: 125 °C, 2,32 bar a; VM: 175 °C, 8,93 bar a; "
         "VA: 250 °C, 39,8 bar a", "Utilidades do Aspen/APEA; h_fg das tabelas de vapor"],
        ["Cargas térmicas", "Tabela 6 (sem integração energética)",
         "Maior demanda de utilidades (partida, ou rede integrada fora de operação)"],
        ["Fluido nas linhas de vapor", "Gás ideal; μ ≈ 1,0 × 10⁻⁵ Pa·s", "P ≤ 1 atm"],
        ["Rugosidade absoluta", "0,015 mm (inox); 0,045 mm (aço carbono)", "White (Moody)"],
        ["NPSH requerido (pré-seleção)", f"{br(NPSH_REQUIRED, 1)} m + margem de "
         f"{br(NPSH_MARGIN, 1)} m", "Bombas centrífugas pequenas; confirmar na curva do fabricante"],
    ]))
    a("")

    # 4 ---------------------------------------------------------------
    a("## 4. Metodologia — método da velocidade (vazão) econômica")
    a("")
    a("### 4.1 Fundamento")
    a("")
    a("O diâmetro econômico minimiza o custo total anual da linha, que soma o custo de "
      "capital anualizado da tubulação, crescente com $D$, e o custo da energia de "
      "bombeamento, que cai rapidamente com $D$:")
    a("")
    a("$$CTA(D) = (a+b)\\,C_{tub}(D) + C_{en}\\,\\frac{Q\\,\\Delta P(D)}{\\eta}\\,t_{op},"
      "\\qquad \\Delta P \\propto \\frac{f\\,L\\,\\rho\\,Q^2}{D^5}$$")
    a("")
    a("Com $C_{tub} \\propto D^{n}$, a condição $dCTA/dD = 0$ leva a "
      "$D_{ót} \\propto Q^{\\approx 0{,}45}\\rho^{\\approx 0{,}13}$ (Peters, Timmerhaus e "
      "West, 2003). Como a vazão ótima é $Q = (\\pi D^2/4)\\,v$, isso equivale a uma "
      "**velocidade econômica** que cresce devagar com o diâmetro. Walas (1988) e Turton "
      "et al. (2018) resumem esse ótimo em heurísticas por tipo de serviço, com o "
      "diâmetro interno $D$ em polegadas:")
    a("")
    a(table(["Serviço", "Velocidade econômica", "ΔP econômico"], [
        ["Líquido — recalque de bomba / linha pressurizada", "$v_{ec} = (5 + D/3)$ ft/s",
         f"2,0 psi/100 ft = {br(2 * PSI_100FT / 10, 1)} kPa/100 m"],
        ["Líquido — sucção de bomba (e líquido saturado)", "$v_{ec} = (1{,}3 + D/6)$ ft/s",
         f"0,4 psi/100 ft = {br(0.4 * PSI_100FT / 10, 2)} kPa/100 m"],
        ["Gás e vapor", "$v_{ec} = 20\\,D$ ft/s",
         f"0,5 psi/100 ft = {br(0.5 * PSI_100FT / 10, 1)} kPa/100 m"],
    ]))
    a("")
    a("### 4.2 Critérios complementares por serviço")
    a("")
    a(table(["Serviço", "Regra econômica", "Restrições adicionais", "Fonte"], [
        ["Recalque, alimentação de colunas, água de resfriamento", "recalque",
         "v ≤ 3 m/s", "Walas; Turton; NBR 5626"],
        ["Sucção de bomba; líquido saturado; condensado antes do purgador; fundo → refervedor",
         "sucção", "NPSH: altura estática mínima (Seção 9.2)", "Walas; Turton"],
        ["Vapor de processo a ≈ 1 atm; vapor d'água", "gás", "v ≤ 40 m/s",
         "Walas; Spirax Sarco (vapor saturado 25–40 m/s)"],
        ["Vapor de processo em vácuo (T-101, T-105)", "gás",
         "ΔP da linha ≤ 1 % da pressão absoluta; v ≤ 60 m/s",
         "Critério adotado: preservar o vácuo e a carga do condensador"],
        ["Condensado do condensador → vaso de refluxo", "—",
         f"Autoventilação: $Fr = v/\\sqrt{{gD}} \\le {mbr(FR_MAX, 2)}$", "Simpson (1968)"],
        ["Bifásico (correntes 6 e 16)", "—",
         "$v \\le V_e = C/\\sqrt{\\rho_m}$ (C = 100, unidades US); v ≥ 3 m/s desejável; "
         "ΔP ≤ 2,0 psi/100 ft", "API RP 14E"],
    ]))
    a("")
    a("### 4.3 Procedimento de seleção")
    a("")
    a("1. Vazão de projeto $Q = 1{,}30\\,Q_n$ (vazão volumétrica nas condições de operação, do Aspen).")
    a("2. Diâmetro econômico $D_{ec}$: solução de $Q = \\dfrac{\\pi D^2}{4}\\,v_{ec}(D)$ "
      "(ou $D_{min}$ de Simpson, ou o diâmetro da velocidade erosional).")
    a("3. Varredura dos diâmetros comerciais da classe, do menor (3/4\") para o maior; "
      "**adota-se o primeiro que atende a todos os critérios**: $v \\le v_{ec}(DI)$, "
      "$\\Delta P/L$ ≤ limite econômico e as restrições da Seção 4.2.")
    a("4. O **critério governante** é a condição que reprovou o diâmetro imediatamente "
      "menor; se o adotado é 3/4\" e $D_{ec}$ < DI(3/4\"), governa o DN mínimo.")
    a("5. Verificações: Reynolds, fator de atrito, ΔP, espessura (B31.3), flambagem em "
      "vácuo, NPSH e comparação com correlações de diâmetro ótimo e com o APEA.")
    a("")
    a("### 4.4 Equações")
    a("")
    a("- Continuidade e Reynolds: $v = \\dfrac{4Q}{\\pi D^2}$, $Re = \\dfrac{\\rho v D}{\\mu}$.")
    a("- Atrito: laminar ($Re$ < 2300) $f = 64/Re$; turbulento, Colebrook-White "
      "$\\dfrac{1}{\\sqrt f} = -2\\log_{10}\\!\\left(\\dfrac{\\varepsilon/D}{3{,}7} + "
      "\\dfrac{2{,}51}{Re\\sqrt f}\\right)$, com Swamee-Jain como estimativa inicial.")
    a("- Darcy-Weisbach: $\\dfrac{\\Delta P}{L} = f\\,\\dfrac{1}{D}\\,\\dfrac{\\rho v^2}{2}$.")
    a("- Vapores (gás ideal): $Q = \\dfrac{\\dot n R T}{P}$, $\\rho = \\dfrac{P\\,M}{R T}$; "
      "vapor de topo $\\dot n_V = D(R+1)$; retorno do refervedor $\\dot n_V$ = boil-up; "
      "alimentação do refervedor kettle = $B + \\dot n_V$ (líquido de fundo).")
    a("- Bifásico homogêneo: $\\rho_m$ e $Q$ do Aspen; "
      "$1/\\mu_m = x/\\mu_V + (1-x)/\\mu_L$ (McAdams).")
    a("- Água de resfriamento: $\\dot m = |\\dot Q|/(c_p\\,\\Delta T)$; vapor: "
      "$\\dot m = \\dot Q/h_{fg}$.")
    a("- NPSH com líquido saturado (vasos de refluxo e fundos de coluna): "
      "$NPSH_d = z - h_f - \\dfrac{K v^2}{2g}$ ⇒ $z_{min} = NPSH_r + \\text{margem} + h_f + "
      "\\dfrac{0{,}5\\,v^2}{2g}$.")
    a("- Espessura (ASME B31.3, eq. 3a): $t = \\dfrac{P D_e}{2(SEW + PY)}$; "
      "$t_{min} = \\dfrac{t + c}{1 - 0{,}125}$ (tolerância de fabricação de 12,5 %).")
    a("- Flambagem elástica de tubo longo sob vácuo: "
      "$P_{cr} = \\dfrac{2E}{1-\\nu^2}\\left(\\dfrac{t}{D_e}\\right)^3$, com fator de segurança 3.")
    a("")

    # 5 ---------------------------------------------------------------
    a("## 5. Propriedades físicas")
    a("")
    a("A massa específica, a massa molar e a vazão volumétrica de cada corrente vêm "
      "diretamente do Aspen (NRTL-HOC). Para as linhas internas das colunas usam-se as "
      "propriedades do destilado (topo) e do produto de fundo. A viscosidade do líquido "
      "foi estimada pela regra de Arrhenius, $\\ln\\mu_m = \\sum x_i \\ln\\mu_i$ (frações "
      "molares), com os componentes puros:")
    a("")
    a(table(["Componente", "Método", "Dado de referência", "μ a 25 °C", "μ a 100 °C"], [
        ["Água", "Equação de Vogel", "—", sig(pure_viscosity("AGUA", 25) * 1e3, 3) + " mPa·s",
         sig(pure_viscosity("AGUA", 100) * 1e3, 3) + " mPa·s"],
        ["Ácido acético", "Lewis–Squires", "1,056 mPa·s a 25 °C (CRC)",
         sig(pure_viscosity("HAC", 25) * 1e3, 3) + " mPa·s",
         sig(pure_viscosity("HAC", 100) * 1e3, 3) + " mPa·s"],
        ["Acetato de benzila", "Lewis–Squires", "≈ 2,0 mPa·s a 25 °C (estimativa)",
         sig(pure_viscosity("BZOAC", 25) * 1e3, 3) + " mPa·s",
         sig(pure_viscosity("BZOAC", 100) * 1e3, 3) + " mPa·s"],
        ["Álcool benzílico", "Andrade, ajustada em 2 pontos",
         "5,47 mPa·s (25 °C) e 2,76 mPa·s (50 °C) (CRC)",
         sig(pure_viscosity("BZOH", 25) * 1e3, 3) + " mPa·s",
         sig(pure_viscosity("BZOH", 100) * 1e3, 3) + " mPa·s"],
    ]))
    a("")
    a("A correlação de Lewis–Squires, $\\mu^{-0{,}2661} = \\mu_{ref}^{-0{,}2661} + "
      "(T - T_{ref})/233$ (μ em cP, T em K; Poling et al., 2001), extrapola a "
      "viscosidade a partir de um único ponto. A incerteza dessas estimativas, de "
      "até ±50 %, é coberta pela análise de sensibilidade (Seção 9.5).")
    a("")

    # 6 ---------------------------------------------------------------
    a("## 6. Seleção de materiais")
    a("")
    a("- **S1 — AISI 316L (ASTM A312 TP316L):** praticamente todas as correntes internas "
      "contêm ácido acético, ou água e acetato de benzila, que hidrolisam formando "
      "ácido. A temperatura chega a 213 °C (fundo da T-103), e o ácido é concentrado "
      "(≈ 90 %) a 112 °C na corrente 29. O aço carbono é inaceitável nesse meio. O "
      "316L (com molibdênio e baixo carbono, adequado à soldagem) é o material usual "
      "para ácido acético e é o mesmo do reator e das colunas T-101, T-103, T-104 e "
      "T-105 (Tabela 29). Este memorial estende a recomendação da Seção 4 do relatório, "
      "que pedia 316L nas correntes 11, 27 e 29, a todas as linhas de processo. Há "
      "ainda dois motivos de processo: o íon ferro, vindo da corrosão do aço carbono, "
      "é trocado pela resina sulfônica DOWEX 50W do R-101 e desativa sítios ácidos; e "
      "ele também prejudica a cor e a pureza do álcool benzílico (99,9 %). O "
      "teor de contaminantes (ácido fórmico, cloretos) deve ser confirmado, porque "
      "eles aceleram a corrosão do 316L em ácido acético quente.")
    a("- **S2 — AISI 304L (ASTM A312 TP304L):** alimentações limpas a 25 °C, sem ácido "
      "(correntes 1, 3 e 4: água de processo e acetato de benzila puro). O 304L é mais "
      "barato que o 316L e mantém a isenção de ferro a montante do catalisador. A "
      "corrente 2 (ácido acético glacial) fica em 316L.")
    a("- **C1 — aço carbono ASTM A106 Gr. B, Classe 150:** água de resfriamento tratada, "
      "vapor de baixa e média pressão e seus condensados; sobre-espessura de corrosão "
      "de 1,5 mm.")
    a("- **C2 — aço carbono ASTM A106 Gr. B, Classe 600:** vapor de alta (≈ 40 bar a, "
      "250 °C) e condensado. A Classe 150 não atende à pressão de projeto de 45 bar g "
      "nessa temperatura.")
    a("")
    a("Schedules: 40S até 1.1/2\" e 10S a partir de 2\" no inox; 80 até 1.1/2\" e 40 a "
      "partir de 2\" no aço carbono, por robustez mecânica dos pequenos diâmetros. "
      "Todos foram verificados pela ASME B31.3 (Seção 9.1). As bombas, que a Tabela 29 "
      "especifica com carcaça de aço carbono, devem acompanhar a classe da linha (316 "
      "nas linhas ácidas).")
    a("")

    # 7 ---------------------------------------------------------------
    a("## 7. Memorial de cálculo — exemplos resolvidos")
    a("")
    a("Os exemplos abaixo mostram o procedimento completo para um caso de cada tipo; "
      "as demais linhas seguem o mesmo cálculo (Seção 8).")
    a("")
    by = {r.line.number: r for r in results}
    a(_example(by["017"], "7.1 Líquido pressurizado — alimentação da T-104 (corrente 17)"))
    a(_example(by["111"], "7.2 Vapor em vácuo — vapor de topo da T-101 (50 torr)"))
    a(_example(by["141"], "7.3 Vapor a 1 atm — vapor de topo da T-104"))
    a(_example(by["215"], "7.4 Água de resfriamento — condensador da T-104"))
    a(_example(by["142"], "7.5 Condensado por gravidade — condensador → vaso de refluxo da T-104"))

    # 8 ---------------------------------------------------------------
    a("## 8. Tabela de cálculo de todas as linhas")
    a("")
    for g in GROUPS:
        a(f"### 8.{GROUPS.index(g) + 1} {g}")
        a("")
        a(_calc_table(by_group[g]))
        a("")
    a(f"![Velocidade × diâmetro]({fig_paths[0]})")
    a("")
    a("*Figura 1 — Velocidade na vazão de projeto de cada linha comparada às curvas de "
      "velocidade econômica. As linhas de líquido de processo ficam abaixo da curva no "
      "DN mínimo, porque a vazão é pequena. As linhas de vapor de maior porte ficam "
      "limitadas pelo teto de 40 m/s ou, em vácuo, pela ΔP admissível.*")
    a("")
    a(f"![Distribuição de DN]({fig_paths[1]})")
    a("")
    a("*Figura 2 — Número de linhas por diâmetro nominal e classe de material.*")
    a("")

    # 9 ---------------------------------------------------------------
    a("## 9. Verificações")
    a("")
    a("### 9.1 Espessura de parede e vácuo")
    a("")
    worst = min(results, key=lambda r: r.wall["t_nom"] / r.wall["t_min"])
    a(f"Todas as {n} linhas atendem à ASME B31.3 com os schedules da Seção 6. A menor "
      f"folga é a da linha `{worst.tag}` ($t_{{min}}$ = {br(worst.wall['t_min'] * 1e3, 2)} mm "
      f"contra $t_{{nom}}$ = {br(worst.wall['t_nom'] * 1e3, 2)} mm). Linhas de vapor sob vácuo:")
    a("")
    vac = [r for r in results if "P_cr" in r.wall and r.line.phase == "V"]
    a(table(["Linha", "NPS / schedule", "De/t", "P_cr/3 [MPa]", "Situação"],
            [[f"`{r.tag}`", f"{r.chosen.size.label} {r.chosen.size.schedule}",
              br(r.chosen.size.OD / r.chosen.size.wall, 0), br(r.wall["P_cr"] / 3e6, 3),
              "✔ ≥ 0,101 MPa"] for r in vac]))
    a("")
    a("### 9.2 NPSH — altura estática mínima nas sucções de líquido saturado")
    a("")
    a("Nos vasos de refluxo e nos fundos das colunas o líquido está no ponto de bolha, "
      "logo $P_{sup} = P_v$ e o NPSH disponível vem só da coluna de líquido. Isso vale "
      "também para a T-101 e a T-105, sob vácuo. A tabela dá a elevação mínima do nível "
      "de líquido sobre o eixo da bomba, a ser respeitada no arranjo (altura de saia das "
      "colunas e do vaso):")
    a("")
    a(table(["Linha", "Descrição", "v [m/s]", "h_f [m]", "z mínimo [m]"],
            [[f"`{r.tag}`", r.line.description, sig(r.chosen.v_d, 2),
              sig(r.dP_line / (r.line.rho * 9.80665), 2), br(r.z_min, 2)]
             for r in results if r.z_min]))
    a("")
    a("### 9.3 Correlações de diâmetro ótimo")
    a("")
    a("Para as linhas fora do DN mínimo, com escoamento turbulento, comparou-se o "
      "resultado com as correlações de Peters, Timmerhaus e West (2003), "
      "$D = 0{,}363\\,Q^{0{,}45}\\rho^{0{,}13}$ (SI), e de Sinnott (2005), "
      "$d = 0{,}260\\,G^{0{,}52}\\rho^{-0{,}37}$ (inox) e $d = 0{,}293\\,G^{0{,}53}\\rho^{-0{,}37}$ "
      "(aço carbono), com d em m e G em kg/s:")
    a("")
    a(_econ_table(results))
    a("")
    a("A correlação de Sinnott fica, na maioria das linhas, a até dois diâmetros "
      "comerciais do resultado. A de Peters–Timmerhaus, deduzida para aço carbono com "
      "custos dos EUA, indica diâmetros maiores, sobretudo em vapor de baixa "
      "densidade. Nenhuma das duas incorpora o DN mínimo, o teto de velocidade ou a "
      "ΔP admissível em vácuo; por isso servem apenas para confirmar a ordem de "
      "grandeza do método da velocidade econômica.")
    a("")
    a("### 9.4 Comparação com os diâmetros do APEA")
    a("")
    a("O relatório do APEA (aba *Stream Pipe Line List*) traz diâmetros de linha "
      "por equipamento, mas **sem vínculo com as correntes** (campos de vazão em "
      "branco). Por isso o APEA usou diâmetros padrão de 20 ou 50 mm, que são "
      "incompatíveis com as vazões reais de vapor:")
    a("")
    a(_apea_table(results))
    a("")
    a("A velocidade que resultaria nos diâmetros do APEA chega a milhares de m/s nas "
      "linhas de vapor de topo e de retorno dos refervedores, o que confirma que o APEA "
      "não dimensionou essas linhas. **Para o P&ID valem os diâmetros deste memorial.** "
      "Isso também significa que o custo de tubulação estimado no APEA está subestimado "
      "nesses itens.")
    a("")
    a("### 9.5 Análise de sensibilidade")
    a("")
    rows = []
    for name, changes in sens:
        txt = ("nenhuma alteração" if not changes else
               "; ".join(f"{num}: {a_} → {b_}" for num, a_, b_ in changes))
        rows.append([name, len(changes), txt])
    a(table(["Caso", "Linhas que mudam de DN", "Alterações (linha: base → caso)"], rows))
    a("")
    a("A seleção é robusta quanto à viscosidade estimada, que pesa pouco porque a maioria "
      "das linhas de líquido é governada pelo DN mínimo. As mudanças aparecem quando se "
      "retira o fator de projeto de 1,30, sobretudo em vapor e água de resfriamento, e "
      "quando se altera o critério de ΔP em vácuo. Esses são os parâmetros a confirmar "
      "com a equipe antes de congelar o P&ID.")
    a("")

    # 10 --------------------------------------------------------------
    a("## 10. Observações sobre o fluxograma e recomendações")
    a("")
    a("1. **P-101 e P-104 são válvulas redutoras de pressão, não bombas.** No Aspen "
      "elas têm ΔP negativa (2 → 0,067 bar e 1,013 → 0,133 bar; Anexos 6 e 12). No "
      "P&ID devem aparecer como válvulas de controle (LV/PV) na alimentação das colunas "
      "a vácuo T-101 e T-105. As linhas 010 e 021 ficam a jusante da válvula; elas "
      "foram mantidas em 3/4\" porque o líquido entra sub-resfriado.")
    a("2. **Falta uma bomba entre M-102 e R-101.** A corrente 6 está a 1,013 bar e "
      "parcialmente vaporizada (0,45 % em massa de vapor), mas o reator opera a 2 bar. "
      "Ela foi dimensionada como bifásica (1\", conservador); com uma bomba de "
      "alimentação do reator, a linha a jusante passa a ser só líquido e cabe em 3/4\".")
    a("3. **Bombas de fundo e de destilado das colunas atmosféricas.** As correntes 14, "
      "15, 18, 19, 28 e 29, e as saídas do decantador V-101 (25 e 26), dependem de "
      "bombas que não aparecem no PFD; o APEA já as inclui. Se elas forem acrescentadas "
      "ao P&ID, a linha de sucção (vaso ou fundo da coluna → bomba) segue o critério "
      "de sucção e a verificação de NPSH da Seção 9.2. O recalque mantém o diâmetro "
      "da corrente na lista de linhas:")
    a("")
    a(table(["Corrente", "Sucção (vaso/fundo → bomba)", "v_proj [m/s]",
             "v_ec [m/s]", "z mínimo [m]", "Recalque (lista de linhas)"],
            [[r.line.number, f"`{r.tag}`", sig(r.chosen.v_d, 2), sig(r.chosen.v_ec, 2),
              br(r.z_min, 2), f"`{by[r.line.number[:3]].tag}`"] for r in extra]))
    a("")
    a("4. **Bocais da T-101 (Tabela 19).** A \"saída topo\" foi calculada com 429 mm "
      "(16,89\"), mas está indicada como 8\" nominal; a \"entrada fundo\" tem 256 mm "
      "(10,08\") e está indicada como 6\". Os valores nominais não correspondem aos "
      "calculados. A linha de vapor de topo resultou em "
      f"{by['111'].chosen.size.label} e o retorno do refervedor em "
      f"{by['116'].chosen.size.label}; os bocais devem ter, no mínimo, o diâmetro da linha.")
    a("5. **Alimentação dos refervedores (linhas 1*k*5).** O líquido que vai ao "
      "refervedor kettle é $B + \\dot n_V$. Sem o perfil de estágios do RadFrac, o "
      "boil-up foi convertido em massa com a massa molar do produto de fundo. Isso "
      "é uma boa aproximação para T-101, T-103 e T-105, cujo vapor de fundo tem "
      "composição próxima à do fundo; é levemente conservador para a T-102 e muito "
      "conservador para a T-104. Nesta, o vapor "
      "de fundo a 115,8 °C é rico em água, e a vazão real deve ser cerca de 3 vezes "
      "menor, o que levaria a linha 145 de "
      f"{by['145'].chosen.size.label} para algo próximo de 3\". Para refinar, "
      "basta preencher `MM_vapor_fundo` (ou `L_fundo_kg_h`) em `dados/colunas.csv` "
      "com o perfil do Aspen (*RadFrac → Profiles*, último estágio) e rodar o "
      "script novamente.")
    a("6. **Condensador da T-103 (206 °C).** A Tabela 6 resfria esse condensador com "
      "água, mas no APEA ele gera vapor de média, e na rede integrada aquece "
      "refervedores. A linha de água 213/214 vale para o caso sem integração. Se o "
      "condensador gerar vapor, substituí-la por linhas de água de caldeira e de vapor.")
    a("7. **Altitude de Paulínia (≈ 590 m, P_atm ≈ 94 kPa).** O Aspen usa 1,013 bar. "
      "Isso não altera as sucções saturadas (Seção 9.2), mas reduz em cerca de 0,7 m "
      "o NPSH de bombas que succionam de tanques abertos para a atmosfera.")
    a("8. **Próximos passos:** isométricos e ΔP real de cada trecho, cálculo de altura "
      "manométrica das bombas, válvulas de controle (ΔP mínima de 10 psi; Walas), "
      "PSVs, linhas do sistema de vácuo, isolamento térmico das linhas acima de 60 °C "
      "(proteção pessoal) e análise de flexibilidade das linhas quentes "
      "(T-103 e vapor de alta).")
    a("")

    # 11 --------------------------------------------------------------
    a("## 11. Referências")
    a("")
    refs = [
        "AMERICAN PETROLEUM INSTITUTE. **API RP 14E**: Recommended practice for design and "
        "installation of offshore production platform piping systems. 5. ed. Washington: API, 1991.",
        "AMERICAN SOCIETY OF MECHANICAL ENGINEERS. **ASME B31.3**: Process piping. New York: ASME.",
        "AMERICAN SOCIETY OF MECHANICAL ENGINEERS. **ASME B36.10M**: Welded and seamless "
        "wrought steel pipe; **ASME B36.19M**: Stainless steel pipe; **ASME B16.5**: Pipe "
        "flanges and flanged fittings. New York: ASME.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. **NBR 5626**: Sistemas prediais de água "
        "fria e água quente — projeto, execução, operação e manutenção. Rio de Janeiro: ABNT, 2020.",
        "HAYNES, W. M. (ed.). **CRC Handbook of Chemistry and Physics**. 97. ed. Boca Raton: "
        "CRC Press, 2016 (viscosidade de líquidos).",
        "PETERS, M. S.; TIMMERHAUS, K. D.; WEST, R. E. **Plant design and economics for "
        "chemical engineers**. 5. ed. New York: McGraw-Hill, 2003.",
        "POLING, B. E.; PRAUSNITZ, J. M.; O'CONNELL, J. P. **The properties of gases and "
        "liquids**. 5. ed. New York: McGraw-Hill, 2001.",
        "SILVA TELLES, P. C. **Tubulações industriais: materiais, projeto, montagem**. "
        "Rio de Janeiro: LTC.",
        "SIMPSON, L. L. Sizing piping for process plants. **Chemical Engineering**, "
        "p. 192-214, 17 jun. 1968.",
        "SINNOTT, R. K. **Coulson & Richardson's Chemical Engineering, v. 6: Chemical "
        "engineering design**. 4. ed. Oxford: Elsevier Butterworth-Heinemann, 2005.",
        "SPIRAX SARCO. **The steam and condensate loop**. Cheltenham: Spirax-Sarco, 2007 "
        "(dimensionamento de linhas de vapor).",
        "TOWLER, G.; SINNOTT, R. **Chemical engineering design: principles, practice and "
        "economics of plant and process design**. 2. ed. Oxford: Butterworth-Heinemann, 2013.",
        "TURTON, R. et al. **Analysis, synthesis and design of chemical processes**. 5. ed. "
        "Upper Saddle River: Prentice Hall, 2018.",
        "WALAS, S. M. **Chemical process equipment: selection and design**. Boston: "
        "Butterworth-Heinemann, 1988 (Rules of thumb — fluid flow).",
        "WHITE, F. M. **Mecânica dos fluidos**. 6. ed. Porto Alegre: AMGH, 2011.",
        "SANTOS, B. C. S.; MENDES, E. M. S.; ALMEIDA, G. A.; PESSOA, H. G. **Desenho de "
        "planta para a produção de álcool benzílico (2,37 t/dia)**. Relatório de Projetos "
        "de Indústrias Químicas I — UNESP, Araraquara, 2026.",
    ]
    for ref in refs:
        a(f"- {ref}")
    a("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def export_docx(md_path, out_path):
    """Word version of the memorial (needs pypandoc-binary); returns success.

    Decimal commas written as {,} inside LaTeX are turned into \\text{} runs so
    Word does not add punctuation spacing after the comma.
    """
    import pathlib
    import re
    import tempfile
    try:
        import pypandoc
    except ImportError:
        return False
    md = pathlib.Path(md_path).read_text(encoding="utf-8")
    md = re.sub(r"(\d+)\{,\}(\d+)", r"\\text{\1,\2}", md)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False,
                                     dir=pathlib.Path(md_path).parent,
                                     encoding="utf-8") as tmp:
        tmp.write(md)
    try:
        pypandoc.convert_file(tmp.name, "docx", outputfile=str(out_path),
                              extra_args=[f"--resource-path={pathlib.Path(md_path).parent}",
                                          "-f", "markdown+tex_math_dollars+pipe_tables"])
    finally:
        pathlib.Path(tmp.name).unlink()
    return True
